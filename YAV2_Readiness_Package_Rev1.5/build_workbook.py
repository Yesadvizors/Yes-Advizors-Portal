#!/usr/bin/env python3
"""build_workbook.py

Deterministic builder for:
    YAV2_Product_Readiness_Matrices_Rev1.5.xlsx

Reads the single authoritative content file (package_content.json) and emits a
workbook of the product-readiness matrices. Workbook metadata is pinned to a
fixed timestamp so the output is a pure function of the input content.
normalize_office.py performs a second deterministic normalisation pass.

Usage: python build_workbook.py [content.json] [output.xlsx]
"""
import datetime
import json
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

HERE = os.path.dirname(os.path.abspath(__file__))
CONTENT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "package_content.json")
OUTPUT = sys.argv[2] if len(sys.argv) > 2 else os.path.join(
    HERE, "YAV2_Product_Readiness_Matrices_Rev1.5.xlsx"
)
FIXED = datetime.datetime(1980, 1, 1, 0, 0, 0)

with open(CONTENT, "r", encoding="utf-8") as fh:
    data = json.load(fh)

HEAD_FILL = PatternFill(start_color="FFDDDDDD", end_color="FFDDDDDD", fill_type="solid")
HEAD_FONT = Font(bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")


def add_sheet(wb, title, headers, rows, widths):
    ws = wb.create_sheet(title=title)
    ws.append(headers)
    for i, _ in enumerate(headers, start=1):
        c = ws.cell(row=1, column=i)
        c.font = HEAD_FONT
        c.fill = HEAD_FILL
        c.alignment = WRAP
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = w
    ws.freeze_panes = "A2"
    return ws


wb = Workbook()
# Remove the default sheet so sheet order is fully controlled.
wb.remove(wb.active)

# 1. Module Catalogue (15)
add_sheet(
    wb, "Module Catalogue",
    ["ID", "Module", "Description", "Readiness"],
    [[m["id"], m["name"], m["description"], m["readiness"]] for m in data["modules"]],
    [10, 30, 60, 26],
)

# 2. UAT Coverage (70)
add_sheet(
    wb, "UAT Coverage",
    ["UAT ID", "Module", "Type", "Title / Scenario", "Expected Result", "Priority"],
    [[c["id"], c["module"], c["type"], c["title"], c["expected"], c["priority"]]
     for c in data["uat_cases"]],
    [10, 10, 14, 55, 45, 10],
)

# 3. UAT Coverage Summary (per module: positive + empty-state present)
modules = data["modules"]
uat = data["uat_cases"]
summary_rows = []
for m in modules:
    mid = m["id"]
    pos = sum(1 for c in uat if c["module"] == mid and c["type"] == "Positive")
    emp = sum(1 for c in uat if c["module"] == mid and c["type"] == "Empty-State")
    tot = sum(1 for c in uat if c["module"] == mid)
    summary_rows.append([
        mid, m["name"], tot,
        "YES" if pos >= 1 else "NO",
        "YES" if emp >= 1 else "NO",
    ])
add_sheet(
    wb, "UAT Coverage Summary",
    ["Module ID", "Module", "Total Cases", "Positive Covered", "Empty-State Covered"],
    summary_rows,
    [12, 30, 12, 18, 20],
)

# 4. Mandatory Validation Traceability (28)
add_sheet(
    wb, "Mandatory Validation",
    ["MV ID", "Validation Rule", "Module", "Field", "UAT Case"],
    [[m["id"], m["rule"], m["module"], m["field"], m["uat_case"]]
     for m in data["mandatory_validations"]],
    [10, 50, 10, 26, 12],
)

# 5. Readiness Consistency
read_rows = [[m["id"], m["name"], m["readiness"]] for m in modules]
add_sheet(
    wb, "Readiness Consistency",
    ["Module ID", "Module", "Readiness Classification"],
    read_rows,
    [12, 32, 28],
)

# 6. SOP Index (15)
add_sheet(
    wb, "SOP Index",
    ["SOP ID", "Module", "Title", "Purpose"],
    [[s["id"], s["module"], s["title"], s["purpose"]] for s in data["sops"]],
    [10, 10, 34, 60],
)

# 7. PJ Decisions (14)
add_sheet(
    wb, "PJ Decisions",
    ["ID", "Topic", "Detail", "Status"],
    [[d["id"], d["topic"], d["detail"], d["status"]] for d in data["pj_decisions"]],
    [8, 42, 52, 10],
)

# 8. Backlog (17)
add_sheet(
    wb, "Backlog",
    ["ID", "Backlog Item", "Priority"],
    [[b["id"], b["title"], b["priority"]] for b in data["backlog"]],
    [8, 60, 12],
)

# 9. Revision History
add_sheet(
    wb, "Revision History",
    ["Revision", "Independent Review Result"],
    [[r["revision"], r["review_result"]] for r in data["revision_history"]],
    [12, 100],
)

# Pin metadata for determinism.
props = wb.properties
props.creator = "Claude.ai (product documentation)"
props.lastModifiedBy = "Claude.ai (product documentation)"
props.created = FIXED
props.modified = FIXED
props.title = "YAV2 Product Readiness Matrices Rev1.5"
props.revision = "1"

wb.save(OUTPUT)
sys.stdout.write(
    "build_workbook.py: wrote %s (%d bytes, %d sheets)\n"
    % (OUTPUT, os.path.getsize(OUTPUT), len(wb.sheetnames))
)
