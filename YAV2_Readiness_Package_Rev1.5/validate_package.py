#!/usr/bin/env python3
"""validate_package.py

Validates the frozen product content and the structural integrity of the built
deliverables, then writes the validation reports. Exits non-zero if ANY check
fails (no partial PASS).

Reports produced:
  * validation_report.md
  * validation_results.json
  * uat_coverage_report.txt
  * mandatory_validation_traceability_report.txt
  * readiness_consistency_report.txt
  * workbook_structural_validation.txt

Usage: python validate_package.py --source-dir .
"""
import json
import os
import sys
import zipfile

from openpyxl import load_workbook

EXPECTED = {
    "modules": 15,
    "uat_cases": 70,
    "positive_coverage": 15,
    "empty_state_coverage": 15,
    "mandatory_validations": 28,
    "sops": 15,
    "pj_decisions": 14,
    "backlog": 17,
}
VALID_READINESS = {"READY", "READY_WITH_CONDITIONS", "NOT_READY", "DOCUMENTATION_ONLY"}
DOCX = "YAV2_Portal_Product_Readiness_UAT_and_SOP_Package_Rev1.5.docx"
XLSX = "YAV2_Product_Readiness_Matrices_Rev1.5.xlsx"


def arg(flag, default=None):
    return sys.argv[sys.argv.index(flag) + 1] if flag in sys.argv else default


def main():
    src = arg("--source-dir", ".")
    checks = []

    def check(name, ok, detail=""):
        checks.append({"check": name, "pass": bool(ok), "detail": detail})
        return ok

    with open(os.path.join(src, "package_content.json"), encoding="utf-8") as fh:
        c = json.load(fh)

    modules = c["modules"]
    uat = c["uat_cases"]
    mv = c["mandatory_validations"]
    sops = c["sops"]
    pj = c["pj_decisions"]
    backlog = c["backlog"]
    mod_ids = {m["id"] for m in modules}
    uat_ids = {u["id"] for u in uat}

    pos = [u for u in uat if u["type"] == "Positive"]
    emp = [u for u in uat if u["type"] == "Empty-State"]
    pos_mods = {u["module"] for u in pos}
    emp_mods = {u["module"] for u in emp}

    # --- content counts ---
    check("modules == 15", len(modules) == EXPECTED["modules"], "found %d" % len(modules))
    check("uat_cases == 70", len(uat) == EXPECTED["uat_cases"], "found %d" % len(uat))
    check("positive coverage 15/15",
          len(pos) == 15 and pos_mods == mod_ids,
          "positive=%d modules=%d" % (len(pos), len(pos_mods)))
    check("empty-state coverage 15/15",
          len(emp) == 15 and emp_mods == mod_ids,
          "empty=%d modules=%d" % (len(emp), len(emp_mods)))
    check("mandatory validations == 28", len(mv) == EXPECTED["mandatory_validations"],
          "found %d" % len(mv))
    check("mandatory validations map to real UAT + module",
          all(m["uat_case"] in uat_ids and m["module"] in mod_ids for m in mv))
    check("sops == 15", len(sops) == EXPECTED["sops"], "found %d" % len(sops))
    check("pj decisions == 14", len(pj) == EXPECTED["pj_decisions"], "found %d" % len(pj))
    check("pj decisions all OPEN", all(d["status"] == "OPEN" for d in pj))
    check("backlog == 17", len(backlog) == EXPECTED["backlog"], "found %d" % len(backlog))
    check("every uat references real module", all(u["module"] in mod_ids for u in uat))
    check("readiness classes valid",
          all(m["readiness"] in VALID_READINESS for m in modules))

    # --- governance / authority freeze ---
    g = c["governance"]
    check("current live authority NONE", g["current_live_authority"] == "NONE")
    check("execution attempts ZERO", g["execution_attempts_authorised"] == "ZERO")
    check("implementation authority NONE", g["implementation_authority"] == "NONE")
    check("V1 access PROHIBITED", g["v1_access"] == "PROHIBITED")
    check("CEP Rev1.6 paused",
          "PAUSED" in g["cep_rev1_6_status"])
    check("package not yet approved",
          "NOT_YET_APPROVED" in g["package_status"])

    # --- DOCX structural ---
    docx_path = os.path.join(src, DOCX)
    docx_ok = os.path.isfile(docx_path)
    docx_parts_ok = False
    if docx_ok:
        z = zipfile.ZipFile(docx_path)
        need = ["[Content_Types].xml", "word/document.xml", "word/styles.xml",
                "docProps/core.xml"]
        docx_parts_ok = all(n in z.namelist() for n in need)
    check("DOCX present", docx_ok, DOCX)
    check("DOCX has required OOXML parts", docx_parts_ok)

    # --- XLSX structural ---
    xlsx_path = os.path.join(src, XLSX)
    xlsx_ok = os.path.isfile(xlsx_path)
    wb_lines = []
    expected_sheets = ["Module Catalogue", "UAT Coverage", "UAT Coverage Summary",
                       "Mandatory Validation", "Readiness Consistency", "SOP Index",
                       "PJ Decisions", "Backlog", "Revision History"]
    sheets_ok = uat_rows_ok = mv_rows_ok = False
    if xlsx_ok:
        wb = load_workbook(xlsx_path, read_only=True)
        sheets_ok = all(s in wb.sheetnames for s in expected_sheets)
        uat_rows = wb["UAT Coverage"].max_row - 1
        mv_rows = wb["Mandatory Validation"].max_row - 1
        mod_rows = wb["Module Catalogue"].max_row - 1
        uat_rows_ok = uat_rows == 70
        mv_rows_ok = mv_rows == 28
        wb_lines.append("Sheets present: %s" % ("YES" if sheets_ok else "NO"))
        wb_lines.append("Module Catalogue data rows: %d (expected 15)" % mod_rows)
        wb_lines.append("UAT Coverage data rows: %d (expected 70)" % uat_rows)
        wb_lines.append("Mandatory Validation data rows: %d (expected 28)" % mv_rows)
        wb.close()
    check("XLSX present", xlsx_ok, XLSX)
    check("XLSX has all expected sheets", sheets_ok)
    check("XLSX UAT Coverage has 70 rows", uat_rows_ok)
    check("XLSX Mandatory Validation has 28 rows", mv_rows_ok)

    overall = all(ck["pass"] for ck in checks)

    # --- write reports ---
    summary = {
        "package": c["package"],
        "revision": c["revision"],
        "overall_pass": overall,
        "expected": EXPECTED,
        "observed": {
            "modules": len(modules), "uat_cases": len(uat),
            "positive_coverage": len(pos), "empty_state_coverage": len(emp),
            "mandatory_validations": len(mv), "sops": len(sops),
            "pj_decisions": len(pj), "pj_decisions_open": sum(1 for d in pj if d["status"] == "OPEN"),
            "backlog": len(backlog),
        },
        "checks": checks,
    }
    with open(os.path.join(src, "validation_results.json"), "w", encoding="utf-8", newline="\n") as fh:
        fh.write(json.dumps(summary, indent=2, sort_keys=True) + "\n")

    md = ["# YAV2 %s — Validation Report" % c["revision"], "",
          "Package: %s" % c["package"], "",
          "Overall result: %s" % ("PASS" if overall else "FAIL"), "",
          "## Content counts", "",
          "| Item | Expected | Observed | Result |", "|---|---|---|---|"]
    rowmap = [("Modules", 15, len(modules)), ("UAT cases", 70, len(uat)),
              ("Positive coverage", 15, len(pos)), ("Empty-state coverage", 15, len(emp)),
              ("Mandatory validations", 28, len(mv)), ("SOPs", 15, len(sops)),
              ("PJ decisions (open)", 14, sum(1 for d in pj if d["status"] == "OPEN")),
              ("Backlog items", 17, len(backlog))]
    for label, exp, obs in rowmap:
        md.append("| %s | %d | %d | %s |" % (label, exp, obs, "PASS" if exp == obs else "FAIL"))
    md += ["", "## All checks", ""]
    for ck in checks:
        md.append("- [%s] %s %s" % ("x" if ck["pass"] else " ", ck["check"],
                                    ("(%s)" % ck["detail"]) if ck["detail"] else ""))
    with open(os.path.join(src, "validation_report.md"), "w", encoding="utf-8", newline="\n") as fh:
        fh.write("\n".join(md) + "\n")

    # uat coverage report
    cov = ["YAV2 %s - UAT COVERAGE REPORT" % c["revision"],
           "Total UAT cases: %d" % len(uat),
           "Positive coverage: %d/15 modules" % len(pos_mods),
           "Empty-state coverage: %d/15 modules" % len(emp_mods), "",
           "%-8s %-30s %-6s %-6s %-6s" % ("Module", "Name", "Total", "Pos", "Empty")]
    for m in modules:
        mid = m["id"]
        t = sum(1 for u in uat if u["module"] == mid)
        p = sum(1 for u in uat if u["module"] == mid and u["type"] == "Positive")
        e = sum(1 for u in uat if u["module"] == mid and u["type"] == "Empty-State")
        cov.append("%-8s %-30s %-6d %-6d %-6d" % (mid, m["name"][:30], t, p, e))
    with open(os.path.join(src, "uat_coverage_report.txt"), "w", encoding="utf-8", newline="\n") as fh:
        fh.write("\n".join(cov) + "\n")

    # mandatory validation traceability report
    tr = ["YAV2 %s - MANDATORY VALIDATION TRACEABILITY REPORT" % c["revision"],
          "Total mandatory validations: %d" % len(mv), "",
          "%-7s %-10s %-26s %s" % ("MV", "Module", "Field", "UAT Case")]
    for m in mv:
        tr.append("%-7s %-10s %-26s %s" % (m["id"], m["module"], m["field"], m["uat_case"]))
    with open(os.path.join(src, "mandatory_validation_traceability_report.txt"), "w",
              encoding="utf-8", newline="\n") as fh:
        fh.write("\n".join(tr) + "\n")

    # readiness consistency report
    rc = ["YAV2 %s - READINESS CONSISTENCY REPORT" % c["revision"],
          "Readiness model: %s" % c["readiness_model"]["overall"],
          "Valid classes: %s" % ", ".join(sorted(VALID_READINESS)), "",
          "%-8s %-32s %s" % ("Module", "Name", "Readiness")]
    for m in modules:
        rc.append("%-8s %-32s %s" % (m["id"], m["name"][:32], m["readiness"]))
    with open(os.path.join(src, "readiness_consistency_report.txt"), "w",
              encoding="utf-8", newline="\n") as fh:
        fh.write("\n".join(rc) + "\n")

    # workbook structural validation report
    wb_report = ["YAV2 %s - WORKBOOK STRUCTURAL VALIDATION" % c["revision"]]
    wb_report += wb_lines if wb_lines else ["XLSX not available for structural validation."]
    wb_report.append("Result: %s" % ("PASS" if (sheets_ok and uat_rows_ok and mv_rows_ok) else "FAIL"))
    with open(os.path.join(src, "workbook_structural_validation.txt"), "w",
              encoding="utf-8", newline="\n") as fh:
        fh.write("\n".join(wb_report) + "\n")

    sys.stdout.write("validate_package.py: overall %s (%d checks)\n"
                     % ("PASS" if overall else "FAIL", len(checks)))
    if not overall:
        for ck in checks:
            if not ck["pass"]:
                sys.stderr.write("  FAIL: %s %s\n" % (ck["check"], ck["detail"]))
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
